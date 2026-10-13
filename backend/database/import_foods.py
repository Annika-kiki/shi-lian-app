"""Idempotent Excel importer. Never connects until an explicit command is run."""
import argparse
import json
import math
from openpyxl import load_workbook
from sqlalchemy import select
from backend.models.food import Food
from backend.services.food import invalidate_catalog

COLUMNS = ('name','aliases','category','kcal_per_100g','protein_per_100g',
           'carbs_per_100g','fat_per_100g','default_quantity_g','unit')

def read_rows(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if 'common_foods_356' not in workbook.sheetnames:
            raise ValueError('缺少工作表common_foods_356')
        iterator = iter(workbook['common_foods_356'].values)
        headers = next(iterator, ())
        if len(set(headers)) != len(headers) or any(c not in headers for c in COLUMNS):
            raise ValueError('列名缺失或重复')
        seen = set()
        for line, values in enumerate(iterator, 2):
            if all(v is None for v in values): continue
            row = {c:values[headers.index(c)] for c in COLUMNS}
            try:
                for c, limit in [('name',64), ('category',32), ('unit',16)]:
                    value = row[c]
                    if value is None and c == 'unit': value = '份'
                    if value is None and c == 'category': continue
                    if not isinstance(value,str) or not value.strip() or len(value.strip()) > limit:
                        raise ValueError(f'{c}无效')
                    row[c] = value.strip()
                if row['aliases'] is not None:
                    if not isinstance(row['aliases'], str): raise ValueError('aliases必须为文本')
                    row['aliases'] = ';'.join(dict.fromkeys(s.strip() for s in row['aliases'].split(';') if s.strip()))
                for c in COLUMNS[3:7]:
                    value = row[c]
                    if value is None and c != 'kcal_per_100g': continue
                    if isinstance(value,bool) or not isinstance(value,(int,float)) or not math.isfinite(value) or not 0 <= value <= 3.4e38:
                        raise ValueError(f'{c}必须为非负有限数值')
                value = row['default_quantity_g']
                if value is None: value = 100
                if isinstance(value,bool) or not isinstance(value,(int,float)) or not math.isfinite(value) or not 0 < value <= 2147483647 or int(value) != value:
                    raise ValueError('默认克数必须为正整数')
                row['default_quantity_g'] = int(value)
                if row['name'] in seen: raise ValueError('Excel食物名称重复')
                seen.add(row['name'])
                yield row, None
            except ValueError as exc:
                yield None, f'第{line}行：{exc}'
    finally:
        workbook.close()

def same_value(existing, incoming):
    # MySQL FLOAT stores single precision; avoid endless updates after round trips.
    if isinstance(existing, (int, float)) and isinstance(incoming, (int, float)):
        return math.isclose(existing, incoming, rel_tol=1e-6, abs_tol=1e-6)
    return existing == incoming

def import_foods(db, path, dry_run=False):
    counts = dict(added=0, updated=0, skipped=0, errors=0)
    messages = []
    for row, error in read_rows(path):
        if error:
            counts['errors'] += 1; messages.append(error); continue
        matches = db.scalars(select(Food).where(Food.name == row['name'])).all()
        if len(matches) > 1:
            counts['errors'] += 1; messages.append(f"数据库名称重复：{row['name']}"); continue
        existing = matches[0] if matches else None
        if not existing:
            counts['added'] += 1
            db.add(Food(**row))
        elif all(same_value(getattr(existing,k), v) for k,v in row.items()): counts['skipped'] += 1
        else:
            counts['updated'] += 1
            for k,v in row.items(): setattr(existing,k,v)
        db.flush()
    if counts['errors'] or dry_run: db.rollback()
    else:
        db.commit(); invalidate_catalog()
    return {**counts, 'committed':not counts['errors'] and not dry_run, 'messages':messages}

def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('xlsx')
    parser.add_argument('--dry-run', action='store_true')
    args = parser.parse_args()
    from backend.database.session import SessionLocal
    with SessionLocal() as db:
        try: result = import_foods(db,args.xlsx,args.dry_run)
        except Exception:
            db.rollback()
            raise
    print(json.dumps(result, ensure_ascii=False))
    return 1 if result['errors'] else 0

if __name__ == '__main__':
    raise SystemExit(main())
